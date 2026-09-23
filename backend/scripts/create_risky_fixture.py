import docx
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent
fixture_dir = ROOT_DIR / "tests" / "fixtures"
fixture_dir.mkdir(parents=True, exist_ok=True)
target_file = fixture_dir / "risky_msme_contract.docx"

doc = docx.Document()
doc.add_heading("ABUSIVE MASTER SERVICES AGREEMENT FOR MSME VENDOR", level=1)

p1 = doc.add_paragraph()
p1.add_run("1. PAYMENT TERMS\n").bold = True
p1.add_run("1.1 Payment shall be made by Purchaser within 120 days of invoice receipt.\n")
p1.add_run("1.2 No interest, penalty, or statutory compensation under the MSMED Act 2006 shall accrue or be payable on delayed payments under any circumstances. Purchaser reserves the right to withhold any payment at its sole discretion without notice.\n")

p2 = doc.add_paragraph()
p2.add_run("2. PENALTY AND DEDUCTIONS\n").bold = True
p2.add_run("2.1 If Supplier delays performance by more than 1 day, Purchaser may deduct 2% per day up to 50% of the contract value directly from pending invoices without prior notice or cure period.\n")

p3 = doc.add_paragraph()
p3.add_run("3. INDEMNIFICATION AND LIABILITY\n").bold = True
p3.add_run("3.1 Supplier shall fully indemnify, defend, and hold harmless Purchaser, its affiliates, officers, and directors against any and all direct, indirect, special, punitive, and consequential losses, lost profits, or third-party claims.\n")
p3.add_run("3.2 Supplier's liability under this Agreement shall be completely uncapped. Purchaser's aggregate liability to Supplier shall not exceed $100.\n")

p4 = doc.add_paragraph()
p4.add_run("4. INTELLECTUAL PROPERTY\n").bold = True
p4.add_run("4.1 All intellectual property, pre-existing tools, algorithms, and deliverables of Supplier shall belong exclusively to Purchaser as work-for-hire upon creation.\n")

p5 = doc.add_paragraph()
p5.add_run("5. TERMINATION\n").bold = True
p5.add_run("5.1 Purchaser may terminate this Agreement immediately for convenience at any time upon 24 hours written notice without paying for unbilled work completed or demobilization expenses.\n")

p6 = doc.add_paragraph()
p6.add_run("6. DISPUTE RESOLUTION\n").bold = True
p6.add_run("6.1 Any dispute arising under this Agreement shall be referred to a sole arbitrator appointed unilaterally by Purchaser.\n")
p6.add_run("6.2 The seat and venue of arbitration shall be London, UK, under ICC rules. Courts in London shall have exclusive jurisdiction. MSME Samadhaan recourse is explicitly waived by Supplier.\n")

doc.save(str(target_file))
print(f"Successfully created risky fixture contract docx at: {target_file}")
